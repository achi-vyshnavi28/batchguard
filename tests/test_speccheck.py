import json
from pathlib import Path

from openpyxl import load_workbook

from speccheck.export import excel_library, jira_csv
from speccheck.generate import draft
from speccheck.lint import lint
from speccheck.parse import parse_prd

PRD = Path(__file__).resolve().parents[1] / "docs" / "prd" / "PRD-1.1_cleaning_log_and_rejection.md"


def test_parse_finds_all_stories_and_criteria():
    title, stories = parse_prd(PRD)
    assert title.startswith("PRD 1.1")
    assert [s.id for s in stories] == ["US-101", "US-102", "US-103", "US-104", "US-105"]
    assert len(stories[0].criteria) == 3 and stories[0].narrative.startswith("As an operator")


def test_lint_flags_the_untestable_requirements():
    _, stories = parse_prd(PRD)
    gaps = lint(stories)
    rules = {(g.story, g.rule) for g in gaps}
    assert ("US-103", "vague_term") in rules            # "quickly", "user-friendly"
    assert ("US-104", "gxp_control_missing") in rules   # override with no reason/signature/audit
    assert ("US-104", "vague_term") in rules            # "as needed"
    assert ("US-105", "placeholder") in rules           # "TBD"
    assert not any(g.story == "US-101" and g.severity == "high" for g in gaps)  # well-specified story


def _fake_llm(prompt: str) -> str:
    return json.dumps({"test_cases": [{"title": "Reject with reason and signature", "type": "positive",
                                       "preconditions": "Batch in review; user qa1", "steps": ["Open batch", "Reject"],
                                       "expected": "Status rejected; audit event", "priority": "high"}],
                       "ambiguities": ["'quickly' has no threshold"]})


def test_versioned_library_keeps_change_history(tmp_path):
    title, stories = parse_prd(PRD)
    drafts = {s.id: draft(s, llm=_fake_llm) for s in stories}
    gaps = lint(stories)
    v1 = excel_library(tmp_path / "test_library_v1.1.xlsx", "1.1", title, stories, drafts, gaps)
    v2 = excel_library(tmp_path / "test_library_v1.2.xlsx", "1.2", title, stories, drafts, gaps, previous=v1)
    wb = load_workbook(v2)
    assert wb.sheetnames[0] == "Change log"
    assert [r[0] for r in wb["Change log"].iter_rows(min_row=2, values_only=True)] == ["1.1", "1.2"]
    assert wb["Test cases"].max_row - 1 == len(stories)


def test_jira_csv_has_stories_and_gap_tasks(tmp_path):
    _, stories = parse_prd(PRD)
    gaps = lint(stories)
    rows = (tmp_path / "j.csv", )
    path = jira_csv(rows[0], stories, gaps, epic="BatchGuard-1.1")
    lines = path.read_text(encoding="utf-8").splitlines()
    assert lines[0] == "Summary,Issue Type,Description,Priority,Labels"
    assert sum(",Story," in line for line in lines) == 5
    assert any("requirement-gap" in line for line in lines)
